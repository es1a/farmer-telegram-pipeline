"""Export aggregated pipeline rows to Excel for officers / demo dashboards."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook

from src.config import DASHBOARD_PATH, DATA_DIR


def append_dashboard_row(row: dict[str, Any], path: Path | None = None) -> Path:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    target = path or DASHBOARD_PATH

    headers = [
        "timestamp",
        "message_id",
        "chat_id",
        "username",
        "farmer_text",
        "topic",
        "urgency",
        "temp_c",
        "condition",
        "soil_moisture_pct",
        "officer_notified",
    ]

    if target.exists():
        from openpyxl import load_workbook

        wb = load_workbook(target)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "pipeline"
        ws.append(headers)

    ws.append(
        [
            row.get("timestamp"),
            row.get("message_id"),
            row.get("chat_id"),
            row.get("username"),
            row.get("farmer_text"),
            row.get("topic"),
            row.get("urgency"),
            row.get("temp_c"),
            row.get("condition"),
            row.get("soil_moisture_pct"),
            row.get("officer_notified"),
        ]
    )
    wb.save(target)
    return target
